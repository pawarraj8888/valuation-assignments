#!/usr/bin/env python3
"""Compare the Excel workbook's calculated values (after a recalculation in Excel) with the
Python package. The inputs are read from the workbook itself, so the check holds for whatever
is typed into the yellow cells. Exit code 1 if any difference exceeds its tolerance.

    ./recalc_excel_mac.sh && python verify_excel.py
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import numpy as np
import openpyxl
from scipy.stats import binom, norm

sys.path.insert(0, str(Path(__file__).resolve().parent))

import build_excel as layout  # noqa: E402
from cdo import BASE_DEAL, N_CASES, SEED, load_fixed_normals, moment_match, simulate, summary_row  # noqa: E402
from cdo.analysis import describe, expected_pool_cash_flows  # noqa: E402
from cdo.defaults import correlate, correlation_matrix  # noqa: E402

N_BONDS, N_PERIODS = layout.N_BONDS, layout.N_PERIODS


def read_block(ws, first_row: int, first_col: int, n_rows: int, n_cols: int) -> np.ndarray:
    rows = ws.iter_rows(min_row=first_row, max_row=first_row + n_rows - 1, min_col=first_col,
                        max_col=first_col + n_cols - 1, values_only=True)
    return np.array(list(rows), dtype=object)


def read_deal(wb) -> dict:
    """The deal as typed into the Inputs sheet."""
    ws = wb["Inputs"]
    deal = dict(BASE_DEAL)
    for cell, key in layout.INPUT_CELLS.values():
        deal[key] = float(ws[cell].value)
    return deal


class Report:
    """Collects one line per check and counts the failures."""

    def __init__(self) -> None:
        self.failures = 0
        print(f"{'check':58s} {'max |Excel - Python|':>22s}  tolerance")

    def compare(self, label: str, excel, python, tolerance: float, relative: bool = False) -> None:
        excel, python = np.asarray(excel, dtype=object), np.asarray(python, dtype=float)
        bad = [v for v in excel.ravel() if not isinstance(v, (int, float)) or isinstance(v, bool)]
        if bad:
            self.failures += 1
            print(f"{label:58s} {'NOT NUMERIC: ' + repr(bad[0]):>22s}  FAIL")
            return
        difference = np.abs(excel.astype(float) - python)
        if relative:
            difference = difference / np.maximum(1.0, np.abs(python))
        worst = float(difference.max())
        ok = worst <= tolerance
        self.failures += 0 if ok else 1
        print(f"{label:58s} {worst:22.3e}  {tolerance:g} {'OK' if ok else 'FAIL'}")


def check_random(report: Report, wb, raw: np.ndarray, Z: np.ndarray) -> None:
    """The Random sheet: stored draws, the moment-matching workings and the matched numbers."""
    ws, first = wb["Random"], layout.FIRST_ROW

    def matrix(column: int) -> np.ndarray:
        return read_block(ws, layout.RAND_MATRIX_FIRST_ROW, column, N_BONDS, N_BONDS)

    def table(column: int) -> np.ndarray:
        return read_block(ws, first, column, N_CASES, N_BONDS)

    demeaned = raw - raw.mean(axis=0)
    covariance = demeaned.T @ demeaned / len(raw)
    factor = np.linalg.cholesky(covariance)
    # Excel keeps 15 significant digits, so a saved value can differ from the csv in the last bit.
    report.compare("Random: fixed normals", table(layout.RAND_RAW_COL), raw, 1e-12)
    report.compare("Random: covariance of the initial numbers", matrix(layout.RAND_RAW_COL), covariance, 1e-12)
    report.compare("Random: Cholesky factor of the covariance", matrix(layout.RAND_DEMEANED_COL), factor, 1e-12)
    report.compare("Random: inverse of the Cholesky factor", matrix(layout.RAND_MATCHED_COL), np.linalg.inv(factor), 1e-10)
    report.compare("Random: de-meaned numbers", table(layout.RAND_DEMEANED_COL), demeaned, 1e-12)
    report.compare("Random: moment-matched numbers", table(layout.RAND_MATCHED_COL), Z, 1e-10)
    report.compare("Random: matched numbers have identity covariance", matrix(layout.RAND_CHECK_COL), np.eye(N_BONDS), 1e-10)
    report.compare("Inputs: number of cases", [wb["Inputs"][layout.CASES_CELL].value], [N_CASES], 0)


def check_model(report: Report, wb, Z: np.ndarray, deal: dict, r: dict) -> None:
    defaults, cash = wb["Defaults"], wb["CashFlows"]
    first = layout.FIRST_ROW
    chol = np.linalg.cholesky(correlation_matrix(N_BONDS, deal["rho"]))
    report.compare("Defaults: Cholesky factor", read_block(defaults, layout.CHOL_FIRST_ROW, 2, N_BONDS, N_BONDS), chol, 1e-12)
    X = correlate(Z, deal["rho"])
    report.compare("Defaults: correlated normals", read_block(defaults, first, layout.DEF_X_COL, N_CASES, N_BONDS), X, 1e-10)
    report.compare("Defaults: uniforms", read_block(defaults, first, layout.DEF_U_COL, N_CASES, N_BONDS), norm.cdf(X), 1e-10)
    if deal["pd"] > 0:
        report.compare("Defaults: default times (relative)", read_block(defaults, first, layout.DEF_T_COL, N_CASES, N_BONDS),
                       r["T"], 1e-9, relative=True)
    report.compare("Defaults: default quarters", read_block(defaults, first, layout.DEF_Q_COL, N_CASES, N_BONDS), r["Q"], 0)
    report.compare("Defaults: number of defaults per case", read_block(defaults, first, layout.DEF_COUNT_COL, N_CASES, 1)[:, 0],
                   r["defaulted"].sum(axis=1), 0)

    blocks = [("pool", layout.CF_POOL_COL, r["pool_cf"]), ("Class A", layout.CF_A_COL, r["a_cf"]),
              ("Class B", layout.CF_B_COL, r["b_cf"]), ("equity", layout.CF_EQ_COL, r["eq_cf"])]
    for k, (label, column, python) in enumerate(blocks):
        report.compare(f"CashFlows: {label}, every case and quarter", read_block(cash, first, column, N_CASES, N_PERIODS), python, 1e-9)
        report.compare(f"CashFlows: {label} total per case", read_block(cash, first, layout.CF_TOTAL_COL + k, N_CASES, 1)[:, 0],
                       python.sum(axis=1), 1e-9)
    for k, (label, due, paid) in enumerate((("A", r["a_due"], r["a_cf"]), ("B", r["b_due"], r["b_cf"]))):
        flags = ((due - paid).sum(axis=1) > 1e-9).astype(int)
        report.compare(f"CashFlows: Class {label} shortfall flag", read_block(cash, first, layout.CF_TOTAL_COL + 4 + k, N_CASES, 1)[:, 0], flags, 0)


def check_statistics(report: Report, wb, Z: np.ndarray, deal: dict, r: dict) -> None:
    ws = wb["Statistics"]
    promised = [r["pool_promised"].sum(), r["a_due"].sum(), r["b_due"].sum()]
    promised.append(promised[0] - promised[1] - promised[2])
    totals = [r[key].sum(axis=1) for key in ("pool_cf", "a_cf", "b_cf", "eq_cf")]
    python = []
    for total, amount in zip(totals, promised):
        d = describe(total, amount)
        python.append([d["no-default amount"], d["mean"], d["std dev"], d["std error"], d["min"], d["5th pct"], d["median"],
                       d["95th pct"], d["max"], d["mean / no-default amount"] if amount else 0.0, (total < amount - 1e-9).mean()])
    report.compare("Statistics: summary table (no-default amount .. share below)", read_block(ws, layout.SUMMARY_FIRST_ROW, 2, 4, 11), python, 1e-9)

    n_defaults = r["defaulted"].sum(axis=1)
    counts = np.bincount(n_defaults, minlength=N_BONDS + 1)
    p_default = 1.0 - (1.0 - deal["pd"]) ** deal["years"]
    report.compare("Statistics: cases by number of defaults", read_block(ws, layout.DIST_FIRST_ROW, 2, N_BONDS + 1, 1)[:, 0], counts, 0)
    report.compare("Statistics: simulated share by number of defaults", read_block(ws, layout.DIST_FIRST_ROW, 3, N_BONDS + 1, 1)[:, 0],
                   counts / N_CASES, 1e-12)
    seen = np.flatnonzero(counts)
    averages = [[totals[0][n_defaults == k].mean(), totals[3][n_defaults == k].mean()] for k in seen]
    report.compare("Statistics: average pool and equity cash by defaults", read_block(ws, layout.DIST_FIRST_ROW, 5, N_BONDS + 1, 2)[seen],
                   averages, 1e-9)
    report.compare("Statistics: binomial benchmark", read_block(ws, layout.DIST_FIRST_ROW, 4, N_BONDS + 1, 1)[:, 0],
                   binom.pmf(np.arange(N_BONDS + 1), N_BONDS, p_default), 1e-12)

    quarterly = np.column_stack([
        r["pool_promised"], expected_pool_cash_flows(deal), r["pool_cf"].mean(axis=0),
        np.percentile(r["pool_cf"], 5, axis=0), np.percentile(r["pool_cf"], 95, axis=0), r["a_cf"].mean(axis=0),
        r["b_cf"].mean(axis=0), r["eq_cf"].mean(axis=0), np.percentile(r["eq_cf"], 5, axis=0),
        np.percentile(r["eq_cf"], 95, axis=0)])
    report.compare("Statistics: quarterly table", read_block(ws, layout.QUARTERLY_FIRST_ROW, 2, N_PERIODS, 10), quarterly, 1e-9)

    histogram = read_block(ws, layout.HIST_FIRST_ROW, 4, layout.HIST_BINS, 2)
    report.compare("Statistics: histogram counts add up to the cases", histogram.sum(axis=0), [N_CASES, N_CASES], 0)

    row = summary_row(Z, deal)
    live = [row[key] for key in ("avg defaults", "P(no default)", "equity mean", "equity std", "equity 5th pct",
                                 "equity min", "P(A shortfall)", "P(B shortfall)")]
    report.compare("Sensitivity: 'This workbook now' row", read_block(wb["Sensitivity"], layout.LIVE_ROW, 3, 1, 8)[0], live, 1e-9)


def check_case(report: Report, wb, raw: np.ndarray, Z: np.ndarray, r: dict) -> None:
    ws = wb["Case"]
    case = int(ws[layout.CASE_CELL].value)
    i = case - 1
    shown = read_block(ws, layout.CASE_BOND_FIRST_ROW, 2, N_BONDS, 2)
    report.compare(f"Case: initial and moment-matched numbers of case {case}", shown, np.column_stack([raw[i], Z[i]]), 1e-10)
    table = read_block(ws, layout.CASE_CF_FIRST_ROW, 2, N_PERIODS, N_BONDS + 4)
    python = np.column_stack([r["bond_cf"][i].T, r["pool_cf"][i], r["a_cf"][i], r["b_cf"][i], r["eq_cf"][i]])
    report.compare(f"Case: cash flows of case {case} (bonds, pool, classes)", table, python, 1e-9)
    quarters = read_block(ws, layout.CASE_BOND_FIRST_ROW, layout.CASE_QUARTER_COL, N_BONDS, 1)[:, 0]
    report.compare(f"Case: default quarters of case {case}", quarters, r["Q"][i], 0)
    check_cell = ws.cell(row=layout.CASE_CF_FIRST_ROW + N_PERIODS + 2, column=12).value
    report.compare("Case: cross-check against the CashFlows sheet", [check_cell], [0.0], 0)


def verify(workbook: Path, data: Path) -> int:
    wb = openpyxl.load_workbook(workbook, data_only=True)
    if wb["CashFlows"].cell(row=layout.FIRST_ROW, column=layout.CF_POOL_COL).value is None:
        print("The workbook has no calculated values yet. Open it in Excel and save it (./recalc_excel_mac.sh), then rerun.")
        return 1
    table, _ = load_fixed_normals(data, N_CASES, N_BONDS, SEED)
    raw = table.values
    Z = moment_match(raw)
    deal = read_deal(wb)
    r = simulate(Z, deal)
    print("Inputs read from the workbook: " + ", ".join(f"{key}={deal[key]:g}" for _, key in layout.INPUT_CELLS.values()))
    report = Report()
    check_random(report, wb, raw, Z)
    check_model(report, wb, Z, deal, r)
    check_statistics(report, wb, Z, deal, r)
    check_case(report, wb, raw, Z, r)
    print(f"\n{'ALL CHECKS PASSED' if report.failures == 0 else f'{report.failures} CHECK(S) FAILED'}")
    return 0 if report.failures == 0 else 1


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", default="../excel/Miniproject3_CDO_Analysis.xlsx")
    parser.add_argument("--data", default="../data/fixed_random_numbers.csv")
    args = parser.parse_args(argv)
    return verify(Path(args.workbook), Path(args.data))


if __name__ == "__main__":
    sys.exit(main())
